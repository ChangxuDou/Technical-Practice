import openpyxl

wb = openpyxl.load_workbook("transactions.xlsx")
print(wb.sheetnames) 

sheet = wb["Sheet1"]

# cell = sheet["a1"]
# print(sheet.max_column)
# print(sheet.max_row)
# print(cell.row)
# print(cell.column)
# print(cell.coordinate)

# read all the elements in the sheet
for row in range(1,sheet.max_row +1):
    for column in range(1,sheet.max_column + 1):
        cell = sheet.cell(row,column)
        print(cell.value)

sheet.append([1,2,3])
wb.save("transaction2.xlsx")


# wb.create_sheet("Sheet2", 0)
# wb.remove(sheet)