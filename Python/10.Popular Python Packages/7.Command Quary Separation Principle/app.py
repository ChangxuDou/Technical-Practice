# CQS: A method should either change state (command)
# or return information without changing state (query).
# Here, sheet.cell() can create missing cells while reading,
# affecting where append() adds the next row.

import openpyxl

wb = openpyxl.load_workbook("transactions.xlsx")
sheet = wb["Sheet1"]
for row in range(1, 10):
    cell = sheet.cell(row, 1)
    print(cell.value)
sheet.append([1, 2, 3])
wb.save('transactions2.xlsx')
